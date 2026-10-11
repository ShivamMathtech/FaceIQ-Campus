import uuid, io
from concurrent.futures import ThreadPoolExecutor
from datetime import timedelta
import numpy as np
import pytest
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError
from openpyxl import load_workbook
from starlette.websockets import WebSocketDisconnect
from conftest import PASSWORD, student_body, course_body, recognize
from app.models.entities import *
from app.computer_vision.engine import engine, match_embedding, quality, decode_image
from app.core.security import cipher


def test_login_invalid_login_logout(env):
    c = env["client"]
    assert (
        c.post(
            "/api/auth/login", json={"email": "admin@example.test", "password": "wrong"}
        ).status_code
        == 401
    )
    r = c.post(
        "/api/auth/login", json={"email": "admin@example.test", "password": PASSWORD}
    )
    assert r.status_code == 200
    h = {"Authorization": "Bearer " + r.json()["access_token"]}
    assert c.get("/api/auth/me", headers=h).json()["role"] == "admin"
    assert c.post("/api/auth/logout", headers=h).status_code == 200
    assert c.get("/api/auth/me", headers=h).status_code == 401


def test_auth_rate_limit(env):
    for _ in range(10):
        env["client"].post(
            "/api/auth/login", json={"email": "bad@example.test", "password": "no"}
        )
    assert (
        env["client"]
        .post("/api/auth/login", json={"email": "bad@example.test", "password": "no"})
        .status_code
        == 429
    )


def test_unique_erp_and_student_edits(env):
    c = env["client"]
    r = c.post("/api/students", json=student_body(), headers=env["admin"])
    assert r.status_code == 201
    assert (
        c.post("/api/students", json=student_body(), headers=env["admin"]).status_code
        == 409
    )
    assert (
        c.patch(
            "/api/students/" + r.json()["id"],
            json={"name": "Updated Name"},
            headers=env["admin"],
        ).json()["name"]
        == "Updated Name"
    )
    assert (
        c.post(
            "/api/students", json=student_body("INVALID"), headers=env["faculty"]
        ).status_code
        == 403
    )


def test_photo_upload_encryption_replacement_withdrawal(workflow):
    w = workflow
    c = w["client"]
    sid = w["st"]["id"]
    h = w["admin"]
    assert c.get("/api/students/" + sid + "/photo", headers=h).status_code == 200
    c.post(
        "/api/students/" + sid + "/photo",
        files={"file": ("face.png", w["image"], "image/png")},
        headers=h,
    )
    with w["factory"]() as db:
        ts = list(
            db.scalars(
                select(BiometricTemplate)
                .where(BiometricTemplate.student_id == sid)
                .order_by(BiometricTemplate.version)
            )
        )
        assert (
            len(ts) == 2
            and not ts[0].active
            and ts[0].encrypted_embedding is None
            and ts[1].active
        )
        assert (
            np.frombuffer(
                cipher().decrypt(ts[1].encrypted_embedding), dtype=np.float32
            )[0]
            == 1
        )
        assert db.get(Student, sid).photo != w["image"]
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 0
    c.delete("/api/students/" + sid + "/biometric-template", headers=h)
    assert c.get("/api/students/" + sid, headers=h).json()["consent"] is False
    assert c.get("/api/students/" + sid + "/photo", headers=h).status_code == 404
    with w["factory"]() as db:
        assert all(
            t.encrypted_embedding is None for t in db.scalars(select(BiometricTemplate))
        )


def test_invalid_image_files_and_missing_consent(env):
    c = env["client"]
    s = c.post("/api/students", json=student_body(), headers=env["admin"]).json()
    assert (
        c.post(
            "/api/students/" + s["id"] + "/photo",
            files={"file": ("fake.png", b"not-image", "image/png")},
            headers=env["admin"],
        ).status_code
        == 422
    )
    assert (
        c.post(
            "/api/students/" + s["id"] + "/photo",
            files={"file": ("fake.txt", b"text", "text/plain")},
            headers=env["admin"],
        ).status_code
        == 422
    )
    c.patch("/api/students/" + s["id"], json={"consent": False}, headers=env["admin"])
    assert (
        c.post(
            "/api/students/" + s["id"] + "/photo",
            files={"file": ("f.png", b"x", "image/png")},
            headers=env["admin"],
        ).status_code
        == 409
    )


def test_course_enrollment_validation_and_scope(workflow):
    w = workflow
    c = w["client"]
    cid = w["course"]["id"]
    assert c.get("/api/courses", headers=w["faculty"]).json()["total"] == 1
    assert c.get("/api/courses", headers=w["other"]).json()["total"] == 0
    assert (
        c.get("/api/courses/" + cid + "/enrollments", headers=w["other"]).status_code
        == 403
    )
    s = c.post(
        "/api/students",
        json={**student_body("WRONG"), "section": "B"},
        headers=w["admin"],
    ).json()
    assert (
        c.post(
            "/api/courses/" + cid + "/enrollments",
            json={"student_ids": [s["id"]]},
            headers=w["admin"],
        ).status_code
        == 422
    )


def test_session_start_authorization_and_overlap(workflow):
    w = workflow
    c = w["client"]
    sid = w["session"]["id"]
    assert (
        c.post("/api/sessions/" + sid + "/start", headers=w["other"]).status_code == 403
    )
    assert (
        c.post("/api/sessions/" + sid + "/end", headers=w["other"]).status_code == 403
    )
    body = {
        k: w["session"][k]
        for k in ("course_id", "faculty_id", "starts_at", "ends_at", "room")
    }
    body["starts_at"] += "Z"
    body["ends_at"] += "Z"
    assert c.post("/api/sessions", json=body, headers=w["admin"]).status_code == 409


def test_face_recognition_records_and_duplicate_is_idempotent(workflow):
    w = workflow
    assert recognize(w).json()["results"][0]["status"] == "PROCESSING"
    rid = str(uuid.uuid4())
    r = recognize(w, rid)
    assert r.status_code == 200, r.text
    assert r.json()["results"][0]["status"] == "PRESENT_RECORDED"
    assert recognize(w).json()["results"][0]["status"] == "ALREADY_PRESENT"
    assert recognize(w, rid).json()["replayed"] is True
    with w["factory"]() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 1
    s = (
        w["client"]
        .get("/api/sessions/" + w["session"]["id"] + "/summary", headers=w["admin"])
        .json()
    )
    assert (
        s["present"] == 1
        and s["attendance_percentage"] == 50
        and s["duplicate_events"] == 1
    )


def test_concurrent_recognition_cannot_duplicate(workflow):
    w = workflow
    recognize(w)
    with ThreadPoolExecutor(max_workers=8) as pool:
        results = list(pool.map(lambda _: recognize(w), range(12)))
    assert all(r.status_code == 200 for r in results), [r.text for r in results]
    statuses = [r.json()["results"][0]["status"] for r in results]
    assert statuses.count("PRESENT_RECORDED") == 1
    with w["factory"]() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 1


def test_database_unique_constraint(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    with w["factory"]() as db:
        db.add(
            AttendanceRecord(
                session_id=w["session"]["id"],
                student_id=w["st"]["id"],
                status="PRESENT",
                source="MANUAL",
            )
        )
        with pytest.raises(IntegrityError):
            db.commit()


def test_unenrolled_not_matched(workflow):
    w = workflow
    w["client"].delete(
        "/api/courses/" + w["course"]["id"] + "/enrollments/" + w["st"]["id"],
        headers=w["admin"],
    )
    assert recognize(w).json()["results"][0]["status"] == "UNKNOWN_FACE"
    with w["factory"]() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 0


def test_internal_boundary_rejects_unenrolled(workflow):
    from app.attendance.service import verify_attendance

    w = workflow
    with w["factory"]() as db:
        outsider = Student(**student_body("OUTSIDER"))
        db.add(outsider)
        db.flush()
        assert (
            verify_attendance(
                db, w["session"]["id"], outsider.id, str(uuid.uuid4()), 0.9
            ).status
            == "NOT_ENROLLED"
        )
        db.commit()


def test_unknown_ambiguous_low_confidence_matching():
    a = np.zeros(128, dtype=np.float32)
    a[0] = 1
    b = np.zeros(128, dtype=np.float32)
    b[1] = 1
    assert match_embedding(a, [("other", b)], 0.55, 0.08)[0] == "UNKNOWN_FACE"
    assert (
        match_embedding(a, [("a", a), ("b", a)], 0.55, 0.08)[0] == "MULTIPLE_CANDIDATES"
    )
    c = a * 0.5 + b * np.sqrt(0.75)
    assert match_embedding(a, [("c", c)], 0.55, 0.08)[0] == "LOW_CONFIDENCE"


def test_ambiguous_pipeline_does_not_record(workflow):
    w = workflow
    w["client"].post(
        "/api/students/" + w["st2"]["id"] + "/photo",
        files={"file": ("face.png", w["image"], "image/png")},
        headers=w["admin"],
    )
    assert recognize(w).json()["results"][0]["status"] == "MULTIPLE_CANDIDATES"
    with w["factory"]() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 0


def test_end_finalizes_absent_blocks_future_and_percentage(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    recognize(w)
    sid = w["session"]["id"]
    assert (
        w["client"]
        .post("/api/sessions/" + sid + "/end", headers=w["faculty"])
        .status_code
        == 200
    )
    assert (
        w["client"]
        .post("/api/sessions/" + sid + "/end", headers=w["faculty"])
        .status_code
        == 200
    )
    assert recognize(w).status_code == 409
    s = w["client"].get("/api/sessions/" + sid + "/summary", headers=w["admin"]).json()
    assert (s["present"], s["absent"], s["pending"], s["attendance_percentage"]) == (
        1,
        1,
        0,
        50,
    )
    assert (
        w["client"]
        .get("/api/dashboard/summary", headers=w["admin"])
        .json()["overall_percentage"]
        == 50
    )


def test_new_session_accepts_student_again(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    w["client"].post(
        "/api/sessions/" + w["session"]["id"] + "/end", headers=w["faculty"]
    )
    # Controlled database clock arrangement makes the old period non-overlapping.
    with w["factory"]() as db:
        old = db.get(AttendanceSession, w["session"]["id"])
        old.starts_at = now() - timedelta(hours=2)
        old.ends_at = now() - timedelta(hours=1)
        db.commit()
    data = dict(
        course_id=w["course"]["id"],
        faculty_id=w["users"][1].id,
        starts_at=(now() - timedelta(minutes=1)).isoformat() + "Z",
        ends_at=(now() + timedelta(minutes=20)).isoformat() + "Z",
        room="Second class",
    )
    w["session"] = (
        w["client"].post("/api/sessions", json=data, headers=w["admin"]).json()
    )
    assert (
        w["client"]
        .post("/api/sessions/" + w["session"]["id"] + "/start", headers=w["faculty"])
        .status_code
        == 200
    )
    recognize(w)
    assert recognize(w).json()["results"][0]["status"] == "PRESENT_RECORDED"
    with w["factory"]() as db:
        assert (
            db.scalar(
                select(func.count())
                .select_from(AttendanceRecord)
                .where(AttendanceRecord.student_id == w["st"]["id"])
            )
            == 2
        )


def test_manual_correction_permissions_and_audit(workflow):
    w = workflow
    url = "/api/sessions/" + w["session"]["id"] + "/manual-attendance"
    body = {
        "student_id": w["st2"]["id"],
        "status": "EXCUSED",
        "reason": "Verified institutional exemption",
    }
    assert w["client"].post(url, json=body, headers=w["faculty"]).status_code == 403
    assert (
        w["client"]
        .post(url, json={**body, "reason": "x"}, headers=w["admin"])
        .status_code
        == 422
    )
    r = w["client"].post(url, json=body, headers=w["admin"])
    assert r.status_code == 200
    r = w["client"].patch(
        "/api/attendance/" + r.json()["id"],
        json={"status": "PRESENT", "reason": "Reviewed attendance evidence"},
        headers=w["admin"],
    )
    assert r.status_code == 200
    logs = (
        w["client"]
        .get("/api/audit-logs?action=ATTENDANCE_CORRECTED", headers=w["admin"])
        .json()["items"]
    )
    assert len(logs) == 2
    assert logs[0]["details"]["before"]["status"] == "EXCUSED"


def test_websocket_authorization_and_logout(workflow):
    w = workflow
    url = "/api/ws/sessions/" + w["session"]["id"]
    origin = {"origin": "http://localhost:5173"}
    with w["client"].websocket_connect(url, headers=origin) as ws:
        ws.send_json({"token": w["tokens"]["one"]})
        data = ws.receive_json()
        assert data["type"] == "snapshot"
        assert data["summary"]["roster_count"] == 2
    with pytest.raises(WebSocketDisconnect):
        with w["client"].websocket_connect(url, headers=origin) as ws:
            ws.send_json({"token": w["tokens"]["two"]})
            ws.receive_json()
    with pytest.raises(WebSocketDisconnect):
        with w["client"].websocket_connect(
            url, headers={"origin": "https://untrusted.invalid"}
        ) as ws:
            ws.send_json({"token": w["tokens"]["admin"]})
            ws.receive_json()


def test_exports_csv_xlsx_and_audit(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    w["client"].post(
        "/api/sessions/" + w["session"]["id"] + "/end", headers=w["faculty"]
    )
    r = w["client"].get("/api/reports/attendance?format=csv", headers=w["faculty"])
    assert r.status_code == 200
    assert "TEST001" in r.text and "TEST002" in r.text
    x = w["client"].get("/api/reports/attendance?format=xlsx", headers=w["faculty"])
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert wb.active.max_row == 3
    assert (
        w["client"]
        .get("/api/audit-logs?action=ATTENDANCE_EXPORTED", headers=w["admin"])
        .json()["total"]
        == 2
    )


def test_student_privacy(workflow):
    w = workflow
    with w["factory"]() as db:
        db.get(Student, w["st"]["id"]).user_id = w["users"][3].id
        db.commit()
    c = w["client"]
    h = w["student"]
    assert c.get("/api/students", headers=h).json()["total"] == 1
    assert c.get("/api/students/" + w["st2"]["id"], headers=h).status_code == 403
    assert (
        c.get("/api/students/" + w["st2"]["id"] + "/photo", headers=h).status_code
        == 403
    )
    assert c.get("/api/audit-logs", headers=h).status_code == 403
    assert (
        c.post("/api/sessions/" + w["session"]["id"] + "/end", headers=h).status_code
        == 403
    )
    c.post("/api/sessions/" + w["session"]["id"] + "/end", headers=w["faculty"])
    rows = c.get("/api/reports/attendance", headers=h).json()["items"]
    assert len(rows) == 1 and rows[0]["erp_id"] == "TEST001"
    assert (
        c.get(
            "/api/sessions/" + w["session"]["id"] + "/recognition-events", headers=h
        ).status_code
        == 403
    )


def test_calibration_and_liveness_fail_closed(workflow):
    w = workflow
    with w["factory"]() as db:
        s = db.get(AttendanceSession, w["session"]["id"])
        s.policy = {**s.policy, "liveness_required": True}
        db.commit()
    assert recognize(w).json()["results"][0]["status"] == "REVIEW_REQUIRED"
    with w["factory"]() as db:
        s = db.get(AttendanceSession, w["session"]["id"])
        s.policy = {
            **s.policy,
            "liveness_required": False,
            "calibration_approved": False,
        }
        db.commit()
    assert recognize(w).json()["results"][0]["status"] == "REVIEW_REQUIRED"


def test_no_model_keeps_management_available(workflow, monkeypatch):
    w = workflow
    monkeypatch.setattr(engine, "detector", None)
    assert (
        w["client"].get("/api/dashboard/summary", headers=w["admin"]).status_code == 200
    )
    assert recognize(w).status_code == 503
    r = w["client"].post(
        "/api/students/" + w["st"]["id"] + "/photo",
        files={"file": ("face.png", w["image"], "image/png")},
        headers=w["admin"],
    )
    assert r.status_code == 200 and r.json()["biometric_status"] == "SETUP_REQUIRED"
    with w["factory"]() as db:
        assert all(not t.active for t in db.scalars(select(BiometricTemplate)))


def test_quality_invalid_dimensions_and_no_face(env, image_bytes, monkeypatch):
    assert quality(np.zeros((100, 100, 3), dtype=np.uint8))
    assert quality(np.full((100, 100, 3), 128, dtype=np.uint8))
    monkeypatch.setattr(engine, "analyze", lambda image: [])
    s = (
        env["client"]
        .post("/api/students", json=student_body(), headers=env["admin"])
        .json()
    )
    r = env["client"].post(
        "/api/students/" + s["id"] + "/photo",
        files={"file": ("face.png", image_bytes, "image/png")},
        headers=env["admin"],
    )
    assert r.status_code == 422


def test_retention_preserves_academic_records(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    with w["factory"]() as db:
        for event in db.scalars(select(RecognitionEvent)):
            event.created_at = now() - timedelta(days=40)
        for t in db.scalars(select(BiometricTemplate)):
            t.created_at = now() - timedelta(days=400)
        db.commit()
    r = w["client"].post("/api/settings/retention/run", headers=w["admin"])
    assert (
        r.status_code == 200
        and r.json()["events_deleted"] == 2
        and r.json()["biometrics_deleted"] == 1
    )
    with w["factory"]() as db:
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 1


def test_csv_formula_injection():
    from app.api.reporting import spreadsheet_safe

    assert spreadsheet_safe("=1+1") == "'=1+1"
    assert spreadsheet_safe("normal") == "normal"


def test_monthly_and_student_analytics_use_unique_records(workflow):
    w = workflow
    recognize(w)
    recognize(w)
    recognize(w)
    w["client"].post(
        "/api/sessions/" + w["session"]["id"] + "/end", headers=w["faculty"]
    )
    data = w["client"].get("/api/analytics/attendance", headers=w["admin"]).json()
    assert data["monthly"][0]["percentage"] == 50
    rows = (
        w["client"].get("/api/analytics/students", headers=w["admin"]).json()["items"]
    )
    assert {r["erp_id"]: r["percentage"] for r in rows} == {
        "TEST001": 100,
        "TEST002": 0,
    }
    assert (
        w["client"].get("/api/analytics/students", headers=w["other"]).json()["total"]
        == 0
    )


def test_recognition_and_finalization_race(workflow):
    w = workflow
    recognize(w)

    def finish():
        return w["client"].post(
            "/api/sessions/" + w["session"]["id"] + "/end", headers=w["faculty"]
        )

    with ThreadPoolExecutor(max_workers=2) as pool:
        requests = [pool.submit(recognize, w), pool.submit(finish)]
        results = [r.result() for r in requests]
    assert results[0].status_code in (200, 409) and results[1].status_code == 200
    with w["factory"]() as db:
        assert db.get(AttendanceSession, w["session"]["id"]).status == "COMPLETED"
        assert db.scalar(select(func.count()).select_from(AttendanceRecord)) == 2


def test_past_window_rejects_recognition(workflow):
    w = workflow
    with w["factory"]() as db:
        s = db.get(AttendanceSession, w["session"]["id"])
        s.starts_at = now() - timedelta(hours=1)
        s.ends_at = now() - timedelta(minutes=1)
        db.commit()
    assert recognize(w).status_code == 409


def test_unknown_sample_resets_temporal_smoothing(workflow, monkeypatch):
    w = workflow
    assert recognize(w).json()["results"][0]["status"] == "PROCESSING"
    analyze = engine.analyze
    monkeypatch.setattr(engine, "analyze", lambda image: [])
    assert recognize(w).json()["results"][0]["status"] == "NO_FACE"
    monkeypatch.setattr(engine, "analyze", analyze)
    assert recognize(w).json()["results"][0]["status"] == "PROCESSING"
    assert recognize(w).json()["results"][0]["status"] == "PRESENT_RECORDED"
